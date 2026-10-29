#!/usr/bin/env python3
"""Generates a realistic sample portfolio workbook for an insurance agency.

The workbook has one sheet per production month (Ekim 2025 - Ekim 2026). The
sheets intentionally differ from each other the way hand-kept agency files do:
header wording, order, upper/lower case, Turkish characters vs. ASCII,
duplicate columns, title rows above the header, TOPLAM rows, blank rows, dates
and amounts stored as text, TC numbers stored as numbers or text, missing or
mistyped TC numbers and two pairs of different customers with the same name.

All people, TC numbers, phone numbers and plates are random and fictional.
The TC numbers satisfy the checksum only so that the validator can be tested.

Outputs (in sample-data/):
  ornek-portfoy.xlsx   the workbook
  expected.json        ground truth used by the automated tests
"""
import datetime as dt
import json
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SEED = 20261029
random.seed(SEED)
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'sample-data')

FIRST_M = ['Ahmet', 'Mehmet', 'Mustafa', 'Ali', 'Hüseyin', 'Hasan', 'İbrahim', 'İsmail', 'Osman', 'Yusuf',
           'Murat', 'Ömer', 'Ramazan', 'Halil', 'Süleyman', 'Abdullah', 'Mahmut', 'Recep', 'Salih', 'Fatih',
           'Kadir', 'Emre', 'Burak', 'Serkan', 'Volkan', 'Cem', 'Can', 'Onur', 'Barış', 'Uğur', 'Erkan', 'Tolga',
           'Kerem', 'Emir', 'Eren', 'Furkan', 'Gökhan', 'Hakan', 'İlker', 'Levent', 'Selim', 'Sinan', 'Tamer',
           'Taner', 'Yasin', 'Zafer', 'Oğuz', 'Doğan', 'Çağlar', 'Şahin']
FIRST_F = ['Ayşe', 'Fatma', 'Emine', 'Hatice', 'Zeynep', 'Elif', 'Meryem', 'Şerife', 'Zehra', 'Sultan', 'Hanife',
           'Merve', 'Özlem', 'Esra', 'Büşra', 'Gül', 'Derya', 'Sevgi', 'Selin', 'Ebru', 'Burcu', 'Aslı', 'Pınar',
           'Gamze', 'Tuğba', 'Ece', 'İrem', 'Dilek', 'Nur', 'Cansu', 'Şule', 'Gizem', 'Yasemin', 'Hülya', 'Songül',
           'Melek', 'Çiğdem', 'Özge', 'Damla', 'Sibel']
LAST = ['Yılmaz', 'Kaya', 'Demir', 'Şahin', 'Çelik', 'Yıldız', 'Yıldırım', 'Öztürk', 'Aydın', 'Özdemir', 'Arslan',
        'Doğan', 'Kılıç', 'Aslan', 'Çetin', 'Kara', 'Koç', 'Kurt', 'Özkan', 'Şimşek', 'Polat', 'Özer', 'Erdoğan',
        'Güneş', 'Aksoy', 'Kaplan', 'Bulut', 'Tekin', 'Karaca', 'Uçar', 'Avcı', 'Korkmaz', 'Turan', 'Ünal',
        'Güler', 'Bozkurt', 'Altun', 'Tunç', 'Sarı', 'Işık', 'Coşkun', 'Bayram', 'Keskin', 'Taş', 'Akın', 'Çakır',
        'Erdem', 'Ateş', 'Gök']
PROFESSIONS = ['Öğretmen', 'Mühendis', 'Doktor', 'Hemşire', 'Esnaf', 'Memur', 'Emekli', 'Avukat', 'Muhasebeci',
               'Şoför', 'Polis', 'Serbest Meslek', 'Ev Hanımı', 'Çiftçi', 'Mimar', 'Eczacı', 'Bankacı',
               'Satış Temsilcisi', 'Teknisyen', 'Akademisyen', 'İşçi', 'Diş Hekimi', 'Esnaf', 'Emekli', 'Memur',
               'Öğretmen', 'Esnaf']
COMPANY_NAMES = ['Yıldız Lojistik Ltd. Şti.', 'Kardeşler Gıda San. Tic. A.Ş.', 'Ege Yapı İnşaat Ltd. Şti.',
                 'Akdeniz Turizm A.Ş.', 'Marmara Tekstil Ltd. Şti.', 'Doğa Eczanesi', 'Özkan Oto Servis',
                 'Bereket Fırını', 'Anadolu Nakliyat Ltd. Şti.', 'Mavi Kırtasiye', 'Güneş Market',
                 'Demir Hukuk Bürosu', 'Çınar Dental Kliniği', 'Atlas Mimarlık', 'Kuzey Elektrik Ltd. Şti.']
SECTORS = ['Lojistik', 'Gıda', 'İnşaat', 'Turizm', 'Tekstil', 'Eczane', 'Oto Servis', 'Fırın', 'Nakliyat',
           'Kırtasiye', 'Market', 'Hukuk', 'Sağlık', 'Mimarlık', 'Elektrik']

# Insurer names, sometimes written differently in different sheets.
INSURERS = {
    'Allianz': ['Allianz Sigorta', 'ALLİANZ', 'Allianz'],
    'Axa': ['Axa Sigorta', 'AXA SİGORTA', 'Axa'],
    'Anadolu': ['Anadolu Sigorta', 'ANADOLU SİGORTA'],
    'Mapfre': ['Mapfre Sigorta', 'MAPFRE'],
    'HDI': ['HDI Sigorta', 'HDI'],
    'Sompo': ['Sompo Sigorta', 'SOMPO'],
    'Türkiye': ['Türkiye Sigorta', 'TÜRKİYE SİGORTA'],
    'Ray': ['Ray Sigorta'],
    'Quick': ['Quick Sigorta', 'QUICK SİGORTA'],
    'Neova': ['Neova Sigorta'],
    'Doğa': ['Doğa Sigorta'],
    'Ankara': ['Ankara Sigorta'],
    'Aksigorta': ['Aksigorta', 'AKSİGORTA'],
    'Unico': ['Unico Sigorta'],
    'Bereket': ['Bereket Sigorta'],
}
INSURER_WEIGHTS = {'Allianz': 14, 'Axa': 12, 'Anadolu': 14, 'Mapfre': 8, 'HDI': 7, 'Sompo': 9, 'Türkiye': 9,
                   'Ray': 5, 'Quick': 6, 'Neova': 4, 'Doğa': 5, 'Ankara': 4, 'Aksigorta': 6, 'Unico': 3,
                   'Bereket': 3}

# kind -> (written names, premium range, commission rate range, vehicle, term days)
POLICY_KINDS = {
    'trafik': (['Trafik', 'TRAFİK', 'Trafik Sigortası', 'Zorunlu Trafik'], (6500, 22000), (0.07, 0.09), True, 365),
    'kasko': (['Kasko', 'KASKO', 'Genişletilmiş Kasko'], (18000, 95000), (0.12, 0.18), True, 365),
    'imm': (['İMM', 'İhtiyari Mali Mesuliyet'], (1200, 5200), (0.14, 0.16), True, 365),
    'dask': (['DASK', 'Dask'], (700, 3600), (0.15, 0.175), False, 365),
    'konut': (['Konut', 'Konut Sigortası', 'KONUT PAKET'], (1500, 7800), (0.20, 0.25), False, 365),
    'tss': (['Tamamlayıcı Sağlık', 'TSS'], (9000, 32000), (0.12, 0.15), False, 365),
    'saglik': (['Özel Sağlık'], (35000, 120000), (0.10, 0.12), False, 365),
    'ferdi': (['Ferdi Kaza'], (400, 3000), (0.20, 0.25), False, 365),
    'seyahat': (['Seyahat Sağlık'], (250, 1800), (0.20, 0.25), False, None),
    'isyeri': (['İşyeri', 'İşyeri Paket'], (4000, 45000), (0.18, 0.22), False, 365),
    'hayat': (['Hayat'], (2000, 12000), (0.15, 0.20), False, 365),
}

TR_MONTHS = ['Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim', 'Kasım',
             'Aralık']
TR_UPPER = {'i': 'İ', 'ı': 'I'}


def tr_upper(s):
    return ''.join(TR_UPPER.get(c, c.upper()) for c in s)


def tr_lower(s):
    return ''.join({'I': 'ı', 'İ': 'i'}.get(c, c.lower()) for c in s)


# ---------------------------------------------------------------- identifiers
def make_tckn():
    d = [random.randint(1, 9)] + [random.randint(0, 9) for _ in range(8)]
    d10 = ((d[0] + d[2] + d[4] + d[6] + d[8]) * 7 - (d[1] + d[3] + d[5] + d[7])) % 10
    d.append(d10)
    d.append(sum(d) % 10)
    return ''.join(map(str, d))


def make_vkn():
    while True:
        d = [random.randint(0, 9) for _ in range(9)]
        total = 0
        for i in range(9):
            tmp = (d[i] + 9 - i) % 10
            v = (tmp * 2 ** (9 - i)) % 9
            if tmp != 0 and v == 0:
                v = 9
            total += v
        check = (10 - total % 10) % 10
        s = ''.join(map(str, d)) + str(check)
        if s[0] != '0':
            return s


def broken_tc(tc):
    """Changes the last digit so the checksum fails (a typo)."""
    return tc[:-1] + str((int(tc[-1]) + 3) % 10)


CITY_CODES = ['34'] * 10 + ['06'] * 4 + ['35'] * 3 + ['16'] * 2 + ['07', '41', '42', '01', '27', '55', '38', '26',
                                                                     '10', '59', '33', '61']
PLATE_LETTERS = 'ABCDEFGHJKLMNPRSTUVYZ'


def make_plate():
    city = random.choice(CITY_CODES)
    n_letters = random.choice([1, 2, 2, 3, 3])
    letters = ''.join(random.choice(PLATE_LETTERS) for _ in range(n_letters))
    digits = random.randint(10, 9999) if n_letters == 1 else random.randint(10, 999)
    if n_letters == 1:
        digits = random.randint(1000, 9999)
    elif n_letters == 2:
        digits = random.randint(100, 9999)
    else:
        digits = random.randint(10, 999)
    return f'{city} {letters} {digits}'


def make_serial():
    return ''.join(random.choice('ABCDEFGHJKLMNPRSTUVYZ') for _ in range(2)) + str(random.randint(100000, 999999))


def make_phone():
    prefix = random.choice(['530', '532', '533', '535', '536', '537', '538', '539', '541', '542', '543', '544', '545',
                            '546', '505', '506', '507', '551', '552', '553', '554', '555'])
    return prefix + ''.join(str(random.randint(0, 9)) for _ in range(7))


def phone_style(p, style):
    if style == 0:
        return f'0{p[:3]} {p[3:6]} {p[6:8]} {p[8:]}'
    if style == 1:
        return p
    if style == 2:
        return f'+90 {p[:3]} {p[3:6]} {p[6:]}'
    if style == 3:
        return f'(0{p[:3]}) {p[3:6]}-{p[6:8]}-{p[8:]}'
    return f'0{p}'


def slug(s):
    table = str.maketrans('çğıöşüÇĞİÖŞÜ', 'cgiosuCGIOSU')
    return s.translate(table).lower().replace(' ', '.')


# ---------------------------------------------------------------- customers
def weighted(d):
    keys = list(d.keys())
    return random.choices(keys, weights=[d[k] for k in keys])[0]


class Customer:
    def __init__(self, cid, name, tc, phone, profession, corporate=False, email=None):
        self.id = cid
        self.name = name
        self.tc = tc
        self.phone = phone
        self.profession = profession
        self.corporate = corporate
        self.email = email
        self.plates = []
        self.family = []


customers = []
used_names = set()
cid = 0


def new_person(name=None):
    global cid
    female = random.random() < 0.42
    if name is None:
        while True:
            name = f'{random.choice(FIRST_F if female else FIRST_M)} {random.choice(LAST)}'
            if name not in used_names:
                break
    used_names.add(name)
    cid += 1
    email = f'{slug(name)}@{random.choice(["gmail.com", "hotmail.com", "outlook.com", "yandex.com"])}' \
        if random.random() < 0.25 else None
    c = Customer(cid, name, make_tckn(), make_phone(), random.choice(PROFESSIONS), email=email)
    # family members that may appear as "Sigortalı adı soyadı" on health policies
    surname = name.split(' ')[-1]
    for _ in range(random.choice([0, 1, 1, 2])):
        c.family.append(f'{random.choice(FIRST_F + FIRST_M)} {surname}')
    customers.append(c)
    return c


for _ in range(300):
    new_person()

# two pairs of different people who share a name (must stay separate customers)
HOMONYMS = ['Mehmet Yılmaz', 'Ayşe Kaya']
for hn in HOMONYMS:
    for _ in range(2):
        cid += 1
        c = Customer(cid, hn, make_tckn(), make_phone(), random.choice(PROFESSIONS))
        customers.append(c)

for i, cname in enumerate(COMPANY_NAMES):
    cid += 1
    c = Customer(cid, cname, make_vkn(), make_phone(), SECTORS[i], corporate=True,
                 email=f'info@{slug(cname.split(" ")[0])}.com.tr')
    customers.append(c)

# ---------------------------------------------------------------- months
MONTHS = [(2025, 10), (2025, 11), (2025, 12)] + [(2026, m) for m in range(1, 11)]
TODAY = dt.date(2026, 10, 29)


def month_key(y, m):
    return f'{y:04d}-{m:02d}'


def random_day_in(y, m):
    first = dt.date(y, m, 1)
    nxt = dt.date(y + (m == 12), m % 12 + 1, 1)
    last_day = (nxt - first).days
    if (y, m) == (2026, 10):
        last_day = min(last_day, 28)  # nothing produced after "today"
    return dt.date(y, m, random.randint(1, last_day))


def add_years(d, n=1):
    try:
        return d.replace(year=d.year + n)
    except ValueError:
        return d.replace(year=d.year + n, day=28)


policy_seq = [100000]


def policy_no(insurer):
    policy_seq[0] += random.randint(1, 97)
    prefix = {'Allianz': 'AZ', 'Axa': 'AX', 'Anadolu': 'AN', 'Mapfre': 'MP', 'HDI': 'HD', 'Sompo': 'SO',
              'Türkiye': 'TS', 'Ray': 'RY', 'Quick': 'QS', 'Neova': 'NV', 'Doğa': 'DG', 'Ankara': 'AK',
              'Aksigorta': 'AS', 'Unico': 'UN', 'Bereket': 'BR'}[insurer]
    return f'{prefix}-{2025 + random.randint(0, 1)}{policy_seq[0]:07d}'


records = []  # ground truth policies


def add_policy(c, kind, ym, plate=None, serial=None, insured=None, start=None, renewal_of=None, new_business=None):
    names, (pmin, pmax), (rmin, rmax), vehicle, term = POLICY_KINDS[kind]
    y, m = ym
    start = start or random_day_in(y, m)
    if term is None:
        end = start + dt.timedelta(days=random.choice([15, 30, 30, 60, 90, 180]))
    else:
        end = add_years(start)
    premium = round(random.uniform(pmin, pmax) / 10) * 10 + random.choice([0, 0, 0.5, 0.25, 0.75])
    if renewal_of is not None:
        premium = round(renewal_of['premium'] * random.uniform(1.25, 1.45), 2)
    commission = round(premium * random.uniform(rmin, rmax), 2)
    insurer = renewal_of['insurer'] if renewal_of is not None and random.random() < 0.7 else weighted(INSURER_WEIGHTS)
    pay = random.random()
    if pay < 0.55:
        card, account, cash = premium, 0, 0
    elif pay < 0.75:
        card, account, cash = 0, 0, premium
    elif pay < 0.9:
        card, account, cash = 0, premium, 0
    else:
        part = round(premium * random.choice([0.3, 0.4, 0.5]), 2)
        card, account, cash = round(premium - part, 2), 0, part
    if new_business is None:
        new_business = random.random() < 0.45
    step = None
    if kind == 'trafik':
        step = random.choices([1, 2, 3, 4, 5, 6, 7, 8], weights=[1, 1, 2, 10, 8, 8, 7, 9])[0]
    elif kind == 'kasko' and random.random() < 0.6:
        step = random.randint(0, 6)
    note = random.choice([''] * 14 + ['Hasarsızlık indirimi uygulandı', '3 taksit', '6 taksit', 'Peşin ödeme',
                                       'Müşteri fiyat karşılaştırdı', 'Kampanyalı fiyat',
                                       'Eski poliçe başka acentede', 'Ruhsat fotoğrafı alındı',
                                       'Ek teminat: İMM', 'Cam muafiyetsiz', 'Yenileme için aranacak'])
    rec = {
        'customer_id': c.id, 'customer_name': c.name, 'tc': c.tc, 'phone': c.phone,
        'profession': c.profession, 'email': c.email,
        'kind': kind, 'policy_name': renewal_of['policy_name'] if renewal_of else random.choice(names),
        'insurer': insurer, 'company_text': random.choice(INSURERS[insurer]),
        'policy_no': policy_no(insurer), 'start': start, 'end': end, 'month': month_key(y, m),
        'plate': plate if vehicle else None, 'serial': serial if vehicle else None,
        'brand': random.choice(['Renault', 'Fiat', 'Toyota', 'Volkswagen', 'Ford', 'Hyundai', 'Honda', 'Opel',
                                'Peugeot', 'Dacia', 'BMW', 'Mercedes']) if vehicle else None,
        'model_year': random.randint(2008, 2026) if vehicle else None,
        'step': step, 'premium': premium, 'commission': commission, 'new_business': new_business,
        'card': card, 'account': account, 'cash': cash, 'note': note,
        'insured': insured or c.name, 'renewal_of': renewal_of['policy_no'] if renewal_of else None,
    }
    records.append(rec)
    return rec


def pick_month(exclude_last=False):
    choices = MONTHS[:-1] if exclude_last else MONTHS
    return random.choice(choices)


for c in customers:
    if c.corporate:
        add_policy(c, 'isyeri', pick_month())
        for _ in range(random.choice([1, 2, 3])):
            plate, serial = make_plate(), make_serial()
            c.plates.append(plate)
            ym = pick_month()
            add_policy(c, 'trafik', ym, plate, serial)
            if random.random() < 0.5:
                add_policy(c, 'kasko', ym, plate, serial)
        continue
    r = random.random()
    if r < 0.62:  # vehicle owner
        for _ in range(1 if random.random() < 0.85 else 2):
            plate, serial = make_plate(), make_serial()
            c.plates.append(plate)
            ym = pick_month()
            if random.random() < 0.88:
                add_policy(c, 'trafik', ym, plate, serial)
            if random.random() < 0.38:
                ym2 = ym if random.random() < 0.5 else pick_month()
                add_policy(c, 'kasko', ym2, plate, serial)
                if random.random() < 0.15:
                    add_policy(c, 'imm', ym2, plate, serial)
    if random.random() < 0.28:  # home
        ym = pick_month()
        add_policy(c, 'dask', ym)
        if random.random() < 0.5:
            add_policy(c, 'konut', ym if random.random() < 0.6 else pick_month())
    if random.random() < 0.16:
        insured = random.choice(c.family) if c.family and random.random() < 0.5 else None
        add_policy(c, 'tss', pick_month(), insured=insured)
    if random.random() < 0.03:
        add_policy(c, 'saglik', pick_month())
    if random.random() < 0.07:
        add_policy(c, 'seyahat', pick_month(), insured=random.choice(c.family) if c.family else None)
    if random.random() < 0.05:
        add_policy(c, 'ferdi', pick_month())
    if random.random() < 0.02:
        add_policy(c, 'hayat', pick_month())

# make sure every person has at least one policy
have = {r['customer_id'] for r in records}
for c in customers:
    if c.id not in have:
        add_policy(c, random.choice(['dask', 'tss', 'ferdi']), pick_month())

# renewals: annual policies written in Ekim 2025 that end in Ekim 2026 before "today"
renewed_count = 0
for rec in list(records):
    if rec['month'] == '2025-10' and POLICY_KINDS[rec['kind']][4] == 365:
        days_left = (rec['end'] - TODAY).days
        if days_left < 0 and random.random() < 0.6:
            c = next(x for x in customers if x.id == rec['customer_id'])
            start = rec['end'] - dt.timedelta(days=random.randint(0, 6))
            if start > TODAY:
                start = TODAY
            add_policy(c, rec['kind'], (2026, 10), rec['plate'], rec['serial'], insured=rec['insured'],
                       start=start, renewal_of=rec, new_business=False)
            renewed_count += 1

# A showcase customer with policies in many different months
showcase = new_person('Kemal Aydoğdu')
showcase.profession = 'Mühendis'
plate_a, serial_a = '34 KMA 482', 'KM482611'
showcase.plates.append(plate_a)
add_policy(showcase, 'trafik', (2025, 11), plate_a, serial_a, new_business=True)
add_policy(showcase, 'kasko', (2026, 1), plate_a, serial_a, new_business=True)
add_policy(showcase, 'dask', (2026, 3), new_business=True)
add_policy(showcase, 'konut', (2026, 3), new_business=True)
add_policy(showcase, 'tss', (2026, 6), insured=f'Elif Aydoğdu', new_business=True)
add_policy(showcase, 'seyahat', (2026, 8), new_business=True)

# a customer with two traffic policies at the same company: a cancellation that names neither car is ambiguous
twin = new_person('Serdar Bozkurt')
twin_recs = []
for ym in [(2026, 2), (2026, 4)]:
    plate, serial = make_plate(), make_serial()
    twin.plates.append(plate)
    r = add_policy(twin, 'trafik', ym, plate, serial, new_business=True)
    r['insurer'], r['company_text'] = 'Sompo', 'Sompo Sigorta'
    r['policy_no'] = policy_no('Sompo')
    r['card'], r['account'], r['cash'] = r['premium'], 0, 0
    twin_recs.append(r)

# ---------------------------------------------------------------- cancellations (İptal block)
# Every month sheet ends with ŞİRKET | POLİÇE | SİGORTALI AD SOYAD | KART | CARİ | NAKİT: policies of an earlier
# month (or the same month) that were cancelled in this month, with the refund per payment type.
BLOCK_HEADERS = ['ŞİRKET', 'POLİÇE', 'SİGORTALI AD SOYAD', 'KART', 'CARİ', 'NAKİT']
renewal_sources = {r['renewal_of'] for r in records if r['renewal_of']}


def month_last_day(ym):
    y, m = ym
    return dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)


def people_key(name):
    return tr_lower(name).strip()


# how many policies of each kind an insured person has: only unambiguous ones are cancelled at random
kind_count = {}
for r in records:
    for n in {people_key(r['insured']), people_key(r['customer_name'])}:
        kind_count[(n, r['kind'])] = kind_count.get((n, r['kind']), 0) + 1


def cancellable(r, ym, same_month=False):
    key = month_key(*ym)
    if same_month:
        if r['month'] != key:
            return False
    elif r['month'] >= key:
        return False
    return (r['kind'] != 'seyahat' and r['customer_name'] not in ('Kemal Aydoğdu', 'Serdar Bozkurt')
            and r['customer_name'] not in HOMONYMS and not r['renewal_of'] and r['policy_no'] not in renewal_sources
            and (r['end'] - TODAY).days > 45 and r['start'] <= month_last_day(ym) and not r.get('cancelled')
            and kind_count[(people_key(r['insured']), r['kind'])] == 1
            and kind_count[(people_key(r['customer_name']), r['kind'])] == 1)


def refund_split(r, ym):
    day = random_day_in(*ym)
    if r['month'] == month_key(*ym):
        day = max(day, r['start'])
    share = max(0.05, min(1.0, (r['end'] - day).days / 365))
    refund = round(r['premium'] * share, 2)
    paid = r['card'] + r['account'] + r['cash']
    card = round(refund * r['card'] / paid, 2) if r['card'] else 0
    account = round(refund * r['account'] / paid, 2) if r['account'] else 0
    cash = round(refund - card - account, 2) if r['cash'] else 0
    return card, account, cash


cancels = []


def cancel_policy(r, ym, policy_text=None, insured=None, expect='auto'):
    card, account, cash = refund_split(r, ym)
    r['cancelled'] = True
    cancels.append({
        'month': month_key(*ym), 'company': random.choice(INSURERS[r['insurer']]),
        'policy': policy_text or random.choice([r['policy_name'], tr_upper(r['policy_name'])]),
        'insured': insured or random.choice([r['insured'], tr_upper(r['insured'])]),
        'card': card, 'account': account, 'cash': cash, 'target': r, 'expect': expect,
    })


for ym in MONTHS[2:]:
    pool = [r for r in records if cancellable(r, ym)]
    for r in random.sample(pool, min(len(pool), random.choice([0, 1, 1, 2, 2, 3]))):
        cancel_policy(r, ym)
# the policy number written in the POLİÇE column
r = random.choice([r for r in records if cancellable(r, (2026, 8))])
cancel_policy(r, (2026, 8), policy_text=r['policy_no'])
# a health policy of a family member
fam = [r for r in records if r['kind'] == 'tss' and r['insured'] != r['customer_name'] and cancellable(r, (2026, 9))]
if fam:
    cancel_policy(fam[0], (2026, 9))
# sold and cancelled in the same month
cancel_policy(random.choice([r for r in records if cancellable(r, (2026, 7), same_month=True)]), (2026, 7))
# two equal candidates: the user picks
cancels.append({'month': '2026-09', 'company': 'SOMPO', 'policy': 'TRAFİK', 'insured': 'SERDAR BOZKURT',
                'card': 1234.5, 'account': 0, 'cash': 0, 'target': None, 'expect': 'ambiguous'})
# a policy from before the file: nothing to find
cancels.append({'month': '2026-05', 'company': 'ANADOLU SİGORTA', 'policy': 'KASKO', 'insured': 'Nuri Tatlı',
                'card': 0, 'account': 0, 'cash': 3200.0, 'target': None, 'expect': 'none'})

# ---------------------------------------------------------------- sheet layouts
V1 = ['Poliçe adı', 'Poliçe bitiş tarihi', 'İletişim', 'Meslek', 'Adı soyadı', 'Tc kimlik no', 'Plaka', 'Seri no',
      'Basamak', 'Poliçe numarası', 'Şirket poliçe tutarı', 'Yeni iş', 'Komisyon', 'Açıklama', 'Şirket',
      'Sigortalı adı soyadı', 'Kart', 'Cari', 'Nakit']
# exactly the list the agency gave, including the duplicates
V2 = ['Poliçe adı', 'Poliçe bitiş tarihi', 'Iletişim', 'Meslek', 'Adi soyadi', 'Police bitis tarihi',
      'Tc kimlik no', 'plaka', 'seri no', 'Meslek', 'Basamak', 'Police numarasi', 'Sirket police tutari', 'Yeni is',
      'Komisyon', 'Aciklama', 'Sirket', 'Sigortali adi soyadi', 'Kart', 'Cari', 'Nakit']
V3 = ['POLİÇE ADI', 'ŞİRKET', 'POLİÇE NUMARASI', 'BAŞLANGIÇ TARİHİ', 'POLİÇE BİTİŞ TARİHİ', 'ADI SOYADI',
      'TC KİMLİK NO', 'İLETİŞİM', 'MESLEK', 'PLAKA', 'SERİ NO', 'MARKA', 'MODEL YILI', 'BASAMAK',
      'ŞİRKET POLİÇE TUTARI', 'KOMİSYON', 'YENİ İŞ', 'KART', 'CARİ', 'NAKİT', 'SİGORTALI ADI SOYADI', 'AÇIKLAMA']
V4 = ['Branş', 'Sigorta Şirketi', 'Poliçe No', 'Vade Sonu', 'Müşteri Adı Soyadı', 'TCKN', 'Telefon', 'Meslek',
      'Plaka', 'Ruhsat Seri No', 'Basamak', 'Brüt Prim', 'Komisyon Tutarı', 'Yeni İş', 'Kredi Kartı', 'Cari',
      'Nakit', 'Sigortalı', 'Açıklama', 'Temsilci', 'E-posta']

YELLOW = PatternFill('solid', fgColor='FFFF00')
BLACK = Side(style='thin', color='000000')
BOX = Border(top=BLACK, bottom=BLACK, left=BLACK, right=BLACK)
BLOCK_ROWS = 17

LAYOUTS = {
    '2025-10': dict(sheet='EKİM 2025', headers=V1, yes=('EVET', 'HAYIR')),
    '2025-11': dict(sheet='KASIM 2025', headers=V1, title='KASIM 2025 ÜRETİM LİSTESİ', yes=('EVET', 'HAYIR')),
    '2025-12': dict(sheet='Aralık 2025', headers=V2, yes=('Evet', '')),
    '2026-01': dict(sheet='OCAK 2026', headers=V1, yes=('EVET', 'HAYIR'), totals=True),
    '2026-02': dict(sheet='ŞUBAT 2026', headers=V3, title='ŞUBAT 2026 - POLİÇE ÜRETİMİ', yes=('YENİ', 'YENİLEME')),
    '2026-03': dict(sheet='MART', headers=V2, yes=('Evet', 'Hayır'), text_values=True),
    '2026-04': dict(sheet='Nisan 2026', headers=V4, yes=('X', ''), block_gap=3),
    '2026-05': dict(sheet='MAYIS 2026', headers=V1, yes=('EVET', 'HAYIR'), totals=True),
    '2026-06': dict(sheet='Haziran 2026', headers=V3, yes=('YENİ', 'YENİLEME'), block_gap=0),
    '2026-07': dict(sheet='Temmuz-2026', headers=V1, yes=('X', ''), blank_rows=True),
    '2026-08': dict(sheet='08.2026', headers=V2, yes=('Evet', '')),
    '2026-09': dict(sheet='EYLÜL 2026', headers=V4, yes=('X', ''), totals=True),
    '2026-10': dict(sheet='Ekim 2026', headers=V1, title='EKİM 2026 ÜRETİM', yes=('EVET', 'HAYIR')),
}

# rows with deliberately missing / mistyped TC numbers (customer still has a name and phone)
by_month = {}
for rec in records:
    by_month.setdefault(rec['month'], []).append(rec)
for k in by_month:
    by_month[k].sort(key=lambda r: (r['start'], r['policy_no']))

individual_recs = [r for r in records if not next(c for c in customers if c.id == r['customer_id']).corporate
                   and r['customer_name'] not in HOMONYMS and r['customer_name'] != 'Kemal Aydoğdu']
# only customers with 2+ rows, so another row still carries the real TC
multi = {}
for r in individual_recs:
    multi.setdefault(r['customer_id'], []).append(r)
multi_ids = [k for k, v in multi.items() if len(v) >= 2]
random.shuffle(multi_ids)
missing_tc_ids = set(multi_ids[:12])
broken_tc_ids = set(multi_ids[12:15])
for cid_ in missing_tc_ids:
    multi[cid_][-1]['tc_override'] = ''
for cid_ in broken_tc_ids:
    multi[cid_][-1]['tc_override'] = broken_tc(multi[cid_][-1]['tc'])

NAME_STYLES = ['title', 'upper', 'title', 'title', 'lower']


def fmt_name(name, style):
    if style == 'upper':
        return tr_upper(name)
    if style == 'lower':
        return tr_lower(name)
    return name


def tr_amount_text(v):
    s = f'{v:,.2f}'  # 12,345.50
    s = s.replace(',', 'X').replace('.', ',').replace('X', '.')
    return random.choice([s, s + ' TL', '₺' + s])


HEADER_FILL = PatternFill('solid', fgColor='1F4E5A')
TITLE_FONT = Font(bold=True, size=14, color='1F4E5A')
THIN = Side(style='thin', color='B7C4C8')

wb = Workbook()
wb.remove(wb.active)
expected_months = {}
sheet_month = {}

for (y, m) in MONTHS:
    key = month_key(y, m)
    lay = LAYOUTS[key]
    ws = wb.create_sheet(lay['sheet'])
    sheet_month[lay['sheet']] = key
    headers = lay['headers']
    row = 1
    if lay.get('title'):
        ws.cell(row=1, column=1, value=lay['title']).font = TITLE_FONT
        ws.cell(row=2, column=1, value=f'Hazırlayan: Merkez Ofis  -  Kayıt tarihi: {random_day_in(y, m):%d.%m.%Y}')
        row = 4
    header_row = row
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=ci, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.column_dimensions[get_column_letter(ci)].width = max(12, min(28, len(h) + 4))
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    # the İptal block at the end of the page, after an empty column (its first column moves with the table)
    bcol = len(headers) + 1 + lay.get('block_gap', 1)
    month_cancels = [c for c in cancels if c['month'] == key]
    for i, h in enumerate(BLOCK_HEADERS):
        cell = ws.cell(row=header_row, column=bcol + i, value=h)
        cell.font = Font(bold=True, color='000000')
        cell.fill = YELLOW
        cell.border = BOX
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(bcol + i)].width = 22 if i == 2 else 13
    for j in range(max(BLOCK_ROWS, len(month_cancels))):
        for i in range(6):
            ws.cell(row=header_row + 1 + j, column=bcol + i).border = BOX
        if j < len(month_cancels):
            c = month_cancels[j]
            vals = [c['company'], c['policy'], c['insured'], c['card'] or None, c['account'] or None, c['cash'] or None]
            for i, v in enumerate(vals):
                cell = ws.cell(row=header_row + 1 + j, column=bcol + i)
                if i >= 3 and v is not None and lay.get('text_values'):
                    cell.value = tr_amount_text(v)
                else:
                    cell.value = v
                    if i >= 3 and v is not None:
                        cell.number_format = '#,##0.00'
    row += 1
    stats = dict(count=0, premium=0.0, commission=0.0, newBusiness=0, card=0.0, account=0.0, cash=0.0)
    text_values = lay.get('text_values', False)
    yes, no = lay['yes']
    for i, rec in enumerate(by_month.get(key, [])):
        if lay.get('blank_rows') and i > 0 and i % 9 == 0:
            row += 1  # an empty spacer row
        name_style = random.choice(NAME_STYLES)
        tc = rec.get('tc_override', rec['tc'])
        phone = phone_style(rec['phone'], random.randint(0, 4))
        contact = phone
        if rec['email'] and random.random() < 0.3 and 'E-posta' not in headers:
            contact = f'{phone} / {rec["email"]}'
        values = {
            'policy_name': rec['policy_name'],
            'end': rec['end'],
            'contact': contact,
            'profession': rec['profession'],
            'name': fmt_name(rec['customer_name'], name_style) + (' ' if random.random() < 0.05 else ''),
            'tc': tc,
            'plate': rec['plate'] or '',
            'serial': rec['serial'] or '',
            'step': rec['step'] if rec['step'] is not None else '',
            'policy_no': rec['policy_no'],
            'premium': rec['premium'],
            'new': yes if rec['new_business'] else no,
            'commission': rec['commission'],
            'note': rec['note'],
            'company': rec['company_text'],
            'insured': fmt_name(rec['insured'], name_style),
            'card': rec['card'] or None,
            'account': rec['account'] or None,
            'cash': rec['cash'] or None,
            'start': rec['start'],
            'brand': rec['brand'] or '',
            'model_year': rec['model_year'] or '',
            'agent': random.choice(['Selin H.', 'Murat K.', 'Ofis']),
            'email': rec['email'] or '',
        }
        header_to_value = {
            'poliçe adı': 'policy_name', 'poliçe bitiş tarihi': 'end', 'police bitis tarihi': 'end',
            'iletişim': 'contact', 'meslek': 'profession', 'adı soyadı': 'name', 'adi soyadi': 'name',
            'tc kimlik no': 'tc', 'plaka': 'plate', 'seri no': 'serial', 'basamak': 'step',
            'poliçe numarası': 'policy_no', 'police numarasi': 'policy_no', 'şirket poliçe tutarı': 'premium',
            'sirket police tutari': 'premium', 'yeni iş': 'new', 'yeni is': 'new', 'komisyon': 'commission',
            'açıklama': 'note', 'aciklama': 'note', 'şirket': 'company', 'sirket': 'company',
            'sigortalı adı soyadı': 'insured', 'sigortali adi soyadi': 'insured', 'kart': 'card', 'cari': 'account',
            'nakit': 'cash', 'başlangıç tarihi': 'start', 'marka': 'brand', 'model yılı': 'model_year',
            'branş': 'policy_name', 'sigorta şirketi': 'company', 'poliçe no': 'policy_no', 'vade sonu': 'end',
            'müşteri adı soyadı': 'name', 'tckn': 'tc', 'telefon': 'contact', 'ruhsat seri no': 'serial',
            'brüt prim': 'premium', 'komisyon tutarı': 'commission', 'kredi kartı': 'card', 'sigortalı': 'insured',
            'temsilci': 'agent', 'e-posta': 'email',
        }
        seen_meslek = False
        for ci, h in enumerate(headers, start=1):
            hk = tr_lower(h)
            if hk == 'ıletişim':
                hk = 'iletişim'
            vkey = header_to_value[hk]
            v = values[vkey]
            if vkey == 'profession':
                if seen_meslek and random.random() < 0.5:
                    v = ''  # second Meslek column often left empty
                seen_meslek = True
            cell = ws.cell(row=row, column=ci)
            if vkey in ('end', 'start'):
                if text_values:
                    cell.value = v.strftime('%d.%m.%Y')
                else:
                    cell.value = v
                    cell.number_format = 'DD.MM.YYYY'
            elif vkey in ('premium', 'commission', 'card', 'account', 'cash'):
                if v is None:
                    cell.value = None
                elif text_values:
                    cell.value = tr_amount_text(v)
                else:
                    cell.value = v
                    cell.number_format = '#,##0.00'
            elif vkey == 'tc':
                if v and headers in (V1, V3) and random.random() < 0.85:
                    cell.value = int(v)
                    cell.number_format = '0'
                elif v and headers is V4 and random.random() < 0.3:
                    cell.value = f'{v[:3]} {v[3:6]} {v[6:9]} {v[9:]}'
                else:
                    cell.value = v
            else:
                cell.value = v if v != '' else None
        stats['count'] += 1
        stats['premium'] += rec['premium']
        stats['commission'] += rec['commission']
        stats['newBusiness'] += 1 if rec['new_business'] else 0
        stats['card'] += rec['card']
        stats['account'] += rec['account']
        stats['cash'] += rec['cash']
        row += 1
    if lay.get('totals'):
        row += 1
        name_col = next(i for i, h in enumerate(headers, start=1) if tr_lower(h) in ('adı soyadı', 'müşteri adı soyadı', 'adi soyadi'))
        ws.cell(row=row, column=name_col, value='TOPLAM').font = Font(bold=True)
        for ci, h in enumerate(headers, start=1):
            hk = tr_lower(h)
            if hk in ('şirket poliçe tutarı', 'brüt prim'):
                ws.cell(row=row, column=ci, value=round(stats['premium'], 2)).font = Font(bold=True)
            elif hk in ('komisyon', 'komisyon tutarı'):
                ws.cell(row=row, column=ci, value=round(stats['commission'], 2)).font = Font(bold=True)
    for k in ('premium', 'commission', 'card', 'account', 'cash'):
        stats[k] = round(stats[k], 2)
    expected_months[key] = stats

os.makedirs(OUT_DIR, exist_ok=True)
xlsx_path = os.path.join(OUT_DIR, 'ornek-portfoy.xlsx')
wb.save(xlsx_path)

# ---------------------------------------------------------------- ground truth
cust_policies = {}
for r in records:
    cust_policies.setdefault(r['customer_id'], []).append(r)
multi_month = [cid_ for cid_, rs in cust_policies.items() if len({r['month'] for r in rs}) >= 2]
show_recs = cust_policies[showcase.id]
renewal_window = []
for r in records:
    days = (r['end'] - TODAY).days
    if 0 <= days <= 30 and r['kind'] != 'seyahat':  # travel policies are not renewed
        renewed = any(x['renewal_of'] == r['policy_no'] for x in records)
        if not renewed:
            renewal_window.append(r['policy_no'])

expected = {
    'seed': SEED,
    'today': TODAY.isoformat(),
    'sheetMonths': sheet_month,
    'months': expected_months,
    'totalPolicies': len(records),
    'totalCustomers': len(cust_policies),
    'multiMonthCustomers': len(multi_month),
    'totalPremium': round(sum(r['premium'] for r in records), 2),
    'totalCommission': round(sum(r['commission'] for r in records), 2),
    'totalAccount': round(sum(r['account'] for r in records), 2),
    'newBusiness': sum(1 for r in records if r['new_business']),
    'renewalsInFile': renewed_count,
    'homonyms': HOMONYMS,
    'missingTcRows': len(missing_tc_ids),
    'brokenTcRows': len(broken_tc_ids),
    'showcase': {
        'name': showcase.name, 'tc': showcase.tc, 'plate': plate_a,
        'months': sorted({r['month'] for r in show_recs}),
        'policyCount': len(show_recs),
        'premium': round(sum(r['premium'] for r in show_recs), 2),
    },
    'dueIn30Days': len(renewal_window),
    'cancellations': {
        'total': len(cancels),
        'auto': sum(1 for c in cancels if c['expect'] == 'auto'),
        'ambiguous': sum(1 for c in cancels if c['expect'] == 'ambiguous'),
        'none': sum(1 for c in cancels if c['expect'] == 'none'),
        'byMonth': {k: sum(1 for c in cancels if c['month'] == k) for k in sorted({c['month'] for c in cancels})},
        'refundTotal': round(sum(c['card'] + c['account'] + c['cash'] for c in cancels), 2),
        'list': [{'month': c['month'], 'insured': c['insured'], 'policy': c['policy'], 'expect': c['expect'],
                  'policyNo': c['target']['policy_no'] if c['target'] else None,
                  'refund': round(c['card'] + c['account'] + c['cash'], 2)} for c in cancels],
    },
    'totalOutstanding': round(sum(max(0.0, r['account'] - sum(c['account'] for c in cancels if c['target'] is r))
                                  for r in records), 2),
    'byKind': {k: sum(1 for r in records if r['kind'] == k) for k in POLICY_KINDS},
    'byInsurer': {k: sum(1 for r in records if r['insurer'] == k) for k in INSURERS},
}
with open(os.path.join(OUT_DIR, 'expected.json'), 'w', encoding='utf-8') as f:
    json.dump(expected, f, ensure_ascii=False, indent=2)

print(f'wrote {xlsx_path}')
print(f'{len(records)} policies, {len(cust_policies)} customers, {len(multi_month)} in 2+ months, '
      f'{renewed_count} renewals inside the file, {len(renewal_window)} due in 30 days, {len(cancels)} cancellations')
